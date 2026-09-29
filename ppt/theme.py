"""Shared look-and-feel for the project deck.

Everything visual lives here: palette, geometry, and the shape helpers that
build_deck.py composes slides out of. Nothing in this module knows about the
project's content -- it only knows how things should look.
"""
import math

from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

# ---------------------------------------------------------------- palette ---
INK      = RGBColor(0x0B, 0x14, 0x24)   # deep navy, dark slides
INK_SOFT = RGBColor(0x14, 0x22, 0x3A)   # raised panel on dark
GHOST    = RGBColor(0x1B, 0x2C, 0x48)   # oversized background numerals
PAPER    = RGBColor(0xFF, 0xFF, 0xFF)
MIST     = RGBColor(0xF3, 0xF6, 0xFA)   # card fill on light slides
LINE     = RGBColor(0xDD, 0xE4, 0xEE)

CYAN  = RGBColor(0x22, 0xD3, 0xEE)
TEAL  = RGBColor(0x0E, 0xA5, 0xA4)
AMBER = RGBColor(0xF5, 0x9E, 0x0B)
ROSE  = RGBColor(0xE1, 0x1D, 0x48)
EMER  = RGBColor(0x10, 0x98, 0x81)
VIOL  = RGBColor(0x7C, 0x3A, 0xED)

TXT      = RGBColor(0x0F, 0x17, 0x2A)   # body on light
TXT_MID  = RGBColor(0x4B, 0x5A, 0x70)
TXT_DIM  = RGBColor(0x8A, 0x97, 0xA8)
ON_DARK  = RGBColor(0xE6, 0xED, 0xF6)
ON_DARK2 = RGBColor(0x9F, 0xB2, 0xCA)

FONT = "Segoe UI"
MONO = "Consolas"
MATH = "Cambria Math"

# --------------------------------------------------------------- geometry ---
W, H = 13.333, 7.5          # inches, 16:9
M = 0.72                    # side margin
BODY_TOP = 1.62             # first usable y on a content slide
BODY_W = W - 2 * M

FOOTER = "Watermark-Guided Tamper Localization and Recovery"


# ---------------------------------------------------------------- helpers ---
def _flat(shape):
    """Kill PowerPoint's default outline and drop-shadow on an autoshape."""
    shape.line.fill.background()
    shape.shadow.inherit = False
    return shape


def rect(slide, x, y, w, h, fill, shape=MSO_SHAPE.RECTANGLE, radius=None):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    _flat(s)
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    s.text_frame.word_wrap = True
    return s


def outline(shape, color, width=1.0):
    shape.line.fill.solid()
    shape.line.fill.fore_color.rgb = color
    shape.line.width = Pt(width)
    return shape


def text(slide, x, y, w, h, body, size=16, color=TXT, bold=False, font=FONT,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, italic=False, spacing=1.0,
         space_after=0, wrap=True):
    """body: str, or list of (text, size, color, bold) / (text, size, color, bold, level)."""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    lines = [body] if isinstance(body, str) else list(body)
    for i, item in enumerate(lines):
        if isinstance(item, str):
            item = (item, size, color, bold)
        t, sz, col, bd = item[:4]
        lvl = item[4] if len(item) > 4 else 0
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.level = lvl
        p.line_spacing = spacing
        p.space_after = Pt(space_after)
        r = p.add_run()
        r.text = t
        r.font.size = Pt(sz)
        r.font.bold = bd
        r.font.italic = italic
        r.font.name = font
        r.font.color.rgb = col
    return box


def bullets(slide, x, y, w, h, items, size=15, color=TXT, gap=7, dot=CYAN,
            lead_color=None):
    """items: list of str, or (str, level). Level 1 renders as a muted sub-line."""
    out = []
    for it in items:
        s, lvl = (it, 0) if isinstance(it, str) else it
        out.append((("▪  " if lvl == 0 else "      –  ") + s,
                    size if lvl == 0 else size - 2,
                    (lead_color or color) if lvl == 0 else TXT_MID,
                    False))
    return text(slide, x, y, w, h, out, spacing=1.18, space_after=gap)


def card(slide, x, y, w, h, fill=MIST, border=LINE, radius=0.045):
    s = rect(slide, x, y, w, h, fill, MSO_SHAPE.ROUNDED_RECTANGLE, radius)
    if border is not None:
        outline(s, border, 1.0)
    return s


def _wrapped_lines(s, size, width_in, bold=False):
    """Rough line count for a string in a box -- enough to stop titles from
    colliding with the body text under them."""
    per_char = size * (0.0082 if bold else 0.0074)   # inches, Segoe UI
    return max(1, math.ceil(len(s) * per_char / max(width_in, 0.1)))


def accent_card(slide, x, y, w, h, title, body, accent=CYAN, fill=MIST,
                title_size=15, body_size=12, on_dark=False, title_h=None):
    """Card with a coloured spine on the left edge."""
    card(slide, x, y, w, h, fill, None if on_dark else LINE)
    rect(slide, x, y, 0.075, h, accent, MSO_SHAPE.ROUNDED_RECTANGLE, 0.5)
    tcol = ON_DARK if on_dark else TXT
    bcol = ON_DARK2 if on_dark else TXT_MID
    if title_h is None:
        n = _wrapped_lines(title, title_size, w - 0.5, bold=True)
        title_h = 0.05 + n * title_size / 60.0
    text(slide, x + 0.28, y + 0.17, w - 0.5, title_h, title, title_size, tcol, True,
         spacing=1.05)
    if body:
        text(slide, x + 0.28, y + 0.17 + title_h, w - 0.5, h - 0.34 - title_h,
             body, body_size, bcol, spacing=1.15)


def stat(slide, x, y, w, h, value, label, sub="", accent=CYAN, on_dark=False,
         value_size=42):
    """Big-number card. The three bands are laid out as fractions of h so the
    value can never land on top of its own label."""
    card(slide, x, y, w, h, INK_SOFT if on_dark else MIST,
         None if on_dark else LINE)
    rect(slide, x, y, w, 0.055, accent)
    text(slide, x + 0.14, y + 0.14 * h, w - 0.28, 0.44 * h, value, value_size,
         accent if on_dark else TXT, True, align=PP_ALIGN.CENTER,
         anchor=MSO_ANCHOR.MIDDLE)
    text(slide, x + 0.14, y + 0.60 * h, w - 0.28, 0.20 * h, label, 12,
         ON_DARK if on_dark else TXT, True, align=PP_ALIGN.CENTER,
         anchor=MSO_ANCHOR.MIDDLE)
    if sub:
        text(slide, x + 0.12, y + 0.79 * h, w - 0.24, 0.19 * h, sub, 9.5,
             ON_DARK2 if on_dark else TXT_MID, align=PP_ALIGN.CENTER,
             anchor=MSO_ANCHOR.MIDDLE, spacing=1.05)


def table(slide, x, y, w, headers, rows, col_w=None, row_h=0.36,
          head_h=0.40, size=11, head_size=10.5, accents=None, zebra=True,
          align_first_left=True):
    """Hand-drawn grid: full control, none of PowerPoint's default banding.

    accents: optional {row_index: RGBColor} to tint one data row.
    """
    n = len(headers)
    col_w = col_w or [w / n] * n
    rect(slide, x, y, w, head_h, INK)
    cx = x
    for c, htxt in enumerate(headers):
        al = PP_ALIGN.LEFT if (c == 0 and align_first_left) else PP_ALIGN.CENTER
        text(slide, cx + 0.12, y + 0.02, col_w[c] - 0.24, head_h, htxt,
             head_size, ON_DARK, True, align=al, anchor=MSO_ANCHOR.MIDDLE)
        cx += col_w[c]
    yy = y + head_h
    for r, row in enumerate(rows):
        acc = (accents or {}).get(r)
        band = MIST if (zebra and r % 2 == 0) else PAPER
        rect(slide, x, yy, w, row_h, band)
        if acc is not None:
            rect(slide, x, yy, 0.055, row_h, acc)
        else:
            rect(slide, x, yy + row_h - 0.012, w, 0.012, LINE)
        cx = x
        for c, val in enumerate(row):
            al = PP_ALIGN.LEFT if (c == 0 and align_first_left) else PP_ALIGN.CENTER
            bold = acc is not None and c == 0
            text(slide, cx + 0.14, yy, col_w[c] - 0.28, row_h, str(val),
                 size, TXT if acc is None else TXT, bold,
                 align=al, anchor=MSO_ANCHOR.MIDDLE)
            cx += col_w[c]
        yy += row_h
    return yy


def chevron(slide, x, y, w, h, label, sub="", fill=INK_SOFT, fg=ON_DARK,
            sub_fg=ON_DARK2, size=12):
    s = rect(slide, x, y, w, h, fill, MSO_SHAPE.CHEVRON)
    s.adjustments[0] = 0.22
    text(slide, x + 0.30, y + (0.16 if sub else 0.0), w - 0.55,
         h - (0.2 if sub else 0.0), label, size, fg, True,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.TOP if sub else MSO_ANCHOR.MIDDLE)
    if sub:
        text(slide, x + 0.30, y + h - 0.44, w - 0.55, 0.36, sub, 9.5, sub_fg,
             align=PP_ALIGN.CENTER, spacing=1.0)
    return s


def arrow(slide, x, y, w, h, color=TXT_DIM, shape=MSO_SHAPE.RIGHT_ARROW):
    s = rect(slide, x, y, w, h, color, shape)
    return s


def formula(slide, x, y, w, h, expr, size=17, color=TXT, align=PP_ALIGN.LEFT):
    return text(slide, x, y, w, h, expr, size, color, font=MATH, align=align,
                anchor=MSO_ANCHOR.MIDDLE)


# ------------------------------------------------------------ slide kinds ---
def blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def content_slide(prs, kicker, title, page=None, dark=False):
    """Standard slide chrome: rule, kicker, title, footer."""
    s = blank(prs)
    bgc = INK if dark else PAPER
    rect(s, 0, 0, W, H, bgc)
    rect(s, 0, 0, W, 0.085, CYAN)
    if kicker:
        text(s, M, 0.44, BODY_W, 0.26, kicker.upper(), 10.5, TEAL if not dark else CYAN,
             True)
    text(s, M, 0.72, BODY_W, 0.62, title, 29, ON_DARK if dark else TXT, True)
    rect(s, M, 1.42, 1.05, 0.045, CYAN)
    # footer
    text(s, M, H - 0.46, 8.0, 0.3, FOOTER, 9, TXT_DIM if not dark else ON_DARK2)
    if page is not None:
        text(s, W - M - 1.2, H - 0.46, 1.2, 0.3, f"{page:02d}", 10,
             TXT_DIM if not dark else ON_DARK2, True, align=PP_ALIGN.RIGHT)
    return s


def section_slide(prs, number, title, blurb=""):
    s = blank(prs)
    rect(s, 0, 0, W, H, INK)
    rect(s, 0, 0, 0.085, H, CYAN)
    # oversized ghost numeral, right side
    text(s, W - 5.0, 0.35, 4.4, 6.0, f"{number:02d}", 260, GHOST, True,
         align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
    text(s, M + 0.35, 2.55, 7.6, 0.4, f"SECTION {number:02d}", 12.5, CYAN, True)
    text(s, M + 0.35, 3.0, 8.4, 1.1, title, 44, ON_DARK, True)
    rect(s, M + 0.35, 4.22, 1.35, 0.05, CYAN)
    if blurb:
        text(s, M + 0.35, 4.52, 7.4, 1.0, blurb, 14, ON_DARK2, spacing=1.25)
    return s
